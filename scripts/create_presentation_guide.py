from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'output/documents/product-configurator-presentation-and-interview-guide.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 10.5)]:
    s = styles[name]; s.font.name = 'Aptos Display' if name != 'Normal' else 'Aptos'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
    s.font.bold = name != 'Normal'; s.paragraph_format.space_before = Pt(14 if name != 'Title' else 0); s.paragraph_format.space_after = Pt(6)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def border(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    for e in ('top','left','bottom','right'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'6'); x.set(qn('w:color'),color); borders.append(x)
    tcPr.append(borders)
def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(9.5)
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, (255,255,255)); shade(t.rows[0].cells[i], '17365D')
        if widths: t.rows[0].cells[i].width=Inches(widths[i])
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,value in enumerate(row):
            set_cell_text(cells[i], value)
            if ridx%2: shade(cells[i], 'EEF3F8')
            if widths: cells[i].width=Inches(widths[i])
    doc.add_paragraph()
    return t
def p(text='', style=None, boldlead=None):
    par=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        par.add_run(boldlead).bold=True; par.add_run(text[len(boldlead):])
    else: par.add_run(text)
    return par
def bullets(items):
    for x in items: p(x, 'List Bullet')
def page(): doc.add_page_break()

# Cover
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Product Configurator Presentation and Interview Guide')
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER; sub.add_run('Custom Logo Canopy Tent  React TypeScript  2D and 3D Product Configuration').italic=True
doc.add_paragraph()
p('Purpose', 'Heading 1')
p('This guide prepares you to present the completed configurator with confidence. It explains the architecture, shows a clear demo flow, provides answers to likely interview questions, and distinguishes what is implemented today from the production extensions you would make next.')
p('Core message', 'Heading 2')
p('The configurator is a reusable product engine. Product data defines options, printable zones, model mappings, and Shopify variants. One structured configuration state drives the 2D editor, 3D preview, pricing, cart payload, and production PDF.')
p('Use this document in two ways:', 'Heading 2')
bullets(['Present Sections 1 through 4 as a 7 to 10 minute project walkthrough.', 'Use Sections 5 through 8 as your technical interview preparation and follow-up reference.'])
page()

p('1  Executive Summary', 'Heading 1')
p('The project is a React and TypeScript product configurator for a custom canopy tent. A customer selects a product size, customizes four printable surfaces with colors, text, and logo uploads, sees the result in a 2D artwork layout and interactive 3D model, receives a dynamic price quote, and can export a production PDF or add the configured product to a mocked Shopify cart.')
table(['Requirement', 'Implemented approach'], [
['Reusable architecture', 'Generic ProductDefinition and ProductConfiguration types keep product data separate from the configurator engine.'],
['2D and 3D synchronization', 'Both views derive from shared Zustand state. The 3D model receives a regenerated canvas texture atlas after design changes.'],
['Dynamic pricing', 'An asynchronous pricing service returns line items, subtotal, GST, total, and a configuration hash.'],
['Shopify readiness', 'Variant resolution, line item properties, and a dynamic customization fee are separated behind a service boundary.'],
['Production handoff', 'A PDF includes selections, artwork details, pricing, Shopify identifiers, and previews.'],
], [1.6,5.7])
p('Opening statement to use', 'Heading 2')
p('“I designed this as a configuration engine rather than a tent-specific page. The customer design exists as structured state, and the editor, 3D preview, pricing, cart handoff, and production PDF are all different consumers of that same state.”')

p('2  Presentation Flow', 'Heading 1')
table(['Time', 'What to show', 'What to say'], [
['0:00 to 0:45', 'Product page and 3D tent', '“The default experience starts with a purchasable product, not an overwhelming design tool. Size selection changes the selected variant, displayed quote, and model asset.”'],
['0:45 to 2:30', 'Open customization studio', '“The tent is split into independently configurable print zones: front, back, left, and right. The design is stored per zone.”'],
['2:30 to 4:00', 'Add text, logo, color; drag and scale', '“Edits write normalized position, rotation, scale, colors, and image references to one central store. The design is data, not merely pixels.”'],
['4:00 to 5:00', 'Toggle or rotate 3D view', '“The 3D view observes the same configuration and regenerates a texture atlas for the GLB fabric material.”'],
['5:00 to 6:00', 'Price summary', '“Pricing is returned by an API-shaped service. Components display the quote; they do not own pricing rules.”'],
['6:00 to 7:00', 'PDF and add to cart', '“The same configuration is converted into a production summary and a Shopify-ready cart request.”'],
], [1.0,1.8,4.5])

p('3  Architecture Walkthrough', 'Heading 1')
p('Start from the entry point.', 'Heading 2')
p('App.tsx mounts Configurator with tentProduct. Configurator is generic; tentProduct is the product-specific definition. This is the first architectural boundary to explain.')
p('Product data', 'Heading 2')
p('data/tentProduct.ts defines the sizes, price minor units, GLB model assets, printable zones, print sizes, and Shopify variant references. Selecting a size affects several downstream systems through data: model asset selection, quote input, and Shopify mapping.')
p('Type system', 'Heading 2')
p('types/product.ts describes the reusable contract. ProductDefinition describes what may be configured. ProductConfiguration describes what the shopper actually configured. A ZoneState stores one surface background plus its ordered layers. A Layer stores text or image data along with normalized x and y coordinates, scale, rotation, and z index.')
p('State management', 'Heading 2')
p('configStore.ts is the source of truth. It exposes deliberate domain actions such as selectOption, addLayer, updateLayer, moveLayerToZone, and resetZone. UI controls call these actions; they do not try to synchronize visual systems directly.')
p('Configuration lifecycle', 'Heading 2')
table(['Stage', 'Input', 'Output'], [
['Customer interaction', 'Size selection or artwork edit', 'A structured update to ProductConfiguration in Zustand.'],
['2D view', 'ZoneState and layers', 'A rendered surface or unfolded tent layout.'],
['3D view', 'All ZoneStates', 'A 2048 by 2048 fabric texture atlas assigned to the GLB.'],
['Pricing', 'ProductDefinition plus ProductConfiguration', 'Async PriceQuote containing authoritative line items and total.'],
['Checkout and production', 'Configuration plus quote', 'Shopify line item properties and a production-summary PDF.'],
], [1.35,2.5,3.45])

p('4  2D and 3D Synchronization', 'Heading 1')
p('The synchronization mechanism is shared state, not manual messaging between two independent editors. When a layer changes, the 2D layout and 3D texture renderer independently react to the same updated ZoneState.')
p('2D path', 'Heading 2')
p('Canvas2DEditor, Toolbar, and LayerList update layers. renderZoneToCanvas draws a zone background and layers in z-index order. Coordinates are normalized, so x 0.5 and y 0.5 mean the center regardless of displayed canvas size.')
p('3D path', 'Heading 2')
p('TentModel loads the selected GLB, clones its scene and fabric material, renders the four zone states into a canopy texture atlas, assigns it to a CanvasTexture, and sets texture.needsUpdate. Cloning prevents multiple configurator instances from mutating shared Three.js material state.')
p('Precise explanation to use', 'Heading 2')
p('“The 2D view and 3D view do not copy state between each other. They are projections of the same configuration. The 2D view shows individual printable surfaces; the 3D renderer maps the full configuration into the UV texture layout of the tent model.”')

p('5  Pricing and Shopify Integration', 'Heading 1')
p('Pricing design', 'Heading 2')
p('pricingService.ts deliberately behaves like a backend service. It accepts the product definition and complete configuration, calculates variant and customization line items, applies GST, and returns a PriceQuote. usePriceQuote debounces requests for 300 milliseconds and ignores stale responses, which prevents old asynchronous results from replacing a newer quote.')
p('Money representation', 'Heading 2')
p('All amounts are stored as integer minor units. For INR, the minor unit is paise: ₹44,900 is stored as 4,490,000. This avoids floating-point arithmetic errors. For a more reusable naming convention, minor-unit names are clearer than names containing Cents when the currency is INR.')
p('Shopify design', 'Heading 2')
p('shopifyService.ts resolves a selected size to a Shopify variant, builds line-item properties for choices and customization details, and models the dynamic customization fee explicitly. This is important because standard Shopify variant prices are fixed while design costs may change per customer.')
table(['Production concern', 'Current design', 'Production replacement'], [
['Quote calculation', 'Mock async getQuote service', 'POST to a pricing API with server-side rules and authorization.'],
['Base product', 'Resolved Shopify variant ID', 'Storefront API cartLinesAdd or theme AJAX cart endpoint.'],
['Dynamic fees', 'customizationFeeCents in request model', 'Cart Transform Shopify Function or an added customization-fee line.'],
['Artwork reference', 'Serialized line-item properties and PDF', 'Upload artwork to controlled storage and save a signed reference URL or order metafield.'],
], [1.55,2.4,3.35])
p('Answer if asked why the UI does not calculate prices', 'Heading 2')
p('“Price is commercial logic and must be trusted. The UI can display a quote, but a real price must come from a server-owned service so rules cannot be changed by a browser client.”')

p('6  PDF and Embedding', 'Heading 1')
p('The PDF is a production handoff artifact, not merely a receipt. pdfService.ts includes product and Shopify references, selections, zone-level design data, layer transforms, price data, and available 2D or 3D previews. This creates a traceable document that could be associated with a Shopify order or fulfillment workflow.')
p('Embedding is handled through an iframe-safe communication boundary. When the cart operation finishes, AddToCartButton sends CONFIGURATOR_ADD_TO_CART through window.postMessage. A host Shopify theme can open its cart drawer or redirect without direct access to the iframe DOM.')

p('7  Interview Questions and Strong Answers', 'Heading 1')
qa = [
('Why Zustand rather than React context or Redux?', '“The configuration is read in distant parts of the tree: editor, 3D preview, pricing, PDF, and cart. Zustand provides a lightweight store with selector-based subscriptions. The provider creates an isolated store per configurator instance, so multiple embedded products do not share state.”'),
('How do you prevent the 2D and 3D views from drifting?', '“There is one ProductConfiguration. The two views do not translate changes between themselves; each renders from the same layer positions, scale, rotation, and zone state. That removes an entire category of synchronization bugs.”'),
('How would you persist unfinished designs?', '“Add a configuration ID and persistence adapter. Save ProductConfiguration to a backend keyed to a customer or anonymous session, debounce saves, version the schema, and restore the configuration before rendering.”'),
('How would you add another product?', '“Keep generic types, state, pricing boundary, cart boundary, and PDF pipeline. Add product data, compatible GLB assets, Shopify mapping, and geometry-specific adapters when the physical shape differs. A pipe needs an unwrapped cylindrical layout rather than the tent quadrant atlas.”'),
('Why are positions normalized?', '“Normalized coordinates are resolution independent. The same design can render on a high-resolution production canvas, compact editor canvas, and GLB texture atlas without storing browser pixel coordinates.”'),
('How would you secure the price?', '“Treat client-side prices as provisional. Recalculate configuration pricing on the backend, validate allowed options and uploaded assets, generate a signed quote or cart payload, and enforce final price with Shopify Functions or server-side cart logic.”'),
('What are the performance risks?', '“Large image uploads, GLB size, texture redraw frequency, and initial bundle size. I would validate image dimensions, compress models with Draco and texture resizing, code-split the 3D studio, throttle or debounce redraws, and profile actual interaction latency.”'),
('What is the limitation of the current 3D mapping?', '“The canopy texture atlas is tailored to supplied tent UVs and fixed face geometry. The generic configuration layer is reusable, but a materially different shape needs its own layout and texture-rendering adapter.”'),
]
for q,a in qa:
    p(q, 'Heading 2'); p(a)

p('8  Deep Dive Follow Up Topics', 'Heading 1')
table(['Topic', 'Points to cover'], [
['Testing', 'Unit test pricing rules and render transforms. Component-test toolbar and layer actions. Add visual regression testing for artwork layout. Use end-to-end tests for configure, quote, PDF, and cart payload flows.'],
['Image quality', 'Read image dimensions at upload, calculate effective DPI against print size and layer scale, warn below a threshold, restrict file type and size, and preserve original production artwork separately from preview assets.'],
['Accessibility', 'Keyboard-operable layer controls, explicit labels, focus management when entering the studio, color contrast, screen-reader summaries for selected surface and price changes.'],
['Reliability', 'Handle GLB, image, quote, and cart failures with visible retries. Track quote version or hash. Persist designs. Use schema migrations as configuration versions evolve.'],
['Security', 'Validate input server-side, scan uploads, store assets privately, issue signed URLs, whitelist option IDs, recompute total server-side, and never trust the browser as the source of final order price.'],
], [1.35,5.95])

p('9  Closing Statement', 'Heading 1')
p('“For the time box, I prioritized the engineering seams that matter in a production configurator: a structured configuration model, a shared source of truth, data-driven product definitions, asynchronous pricing, explicit Shopify responsibilities, and a production handoff artifact. The mock services are intentionally replaceable so that a real integration can be introduced without rewriting the configurator UI.”')
p('Pre interview checklist', 'Heading 2')
bullets(['Run npm run build and npm run lint before the interview.', 'Prepare one small logo image for the live customization demo.', 'Keep the README open at the architecture diagram and known limitations section.', 'Practice the seven-minute walkthrough once without reading this document.', 'Be clear about scope: this is a strong time-boxed architecture demo, not a claim of finished production infrastructure.'])

for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('Product Configurator Presentation and Interview Guide'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
